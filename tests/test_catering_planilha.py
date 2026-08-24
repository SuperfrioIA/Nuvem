"""V3.3 -- planilha aberta, download e auditoria.

## O aceite deste lote

`test_planilha_somada_bate_com_a_matriz`: somando **todas as paginas** da
planilha, o total tem que dar exatamente o que a Matriz agrega no mesmo
recorte. As duas leem o mesmo `recorte.de_para_where()`, entao o teste nao esta
provando que eu escrevi o mesmo `WHERE` duas vezes -- esta provando que
agregacao e detalhe **contam as mesmas linhas**, que e onde erro de
paginacao, de `LIMIT/OFFSET` e de `JOIN` duplicando linha apareceriam.

`test_download_bate_com_o_csv_de_origem`: o arquivo baixado somado contra o CSV
do DW, no mesmo recorte. Mesma ideia do aceite do V3.2, agora no caminho do
download.

## O que este arquivo guarda de verdade

- **paginacao deterministica**: sem ordem total, pagina 2 pode repetir linha da
  pagina 1 e omitir outra, sem erro nenhum;
- **zero a esquerda no xlsx**: `0000000609` tem que sair como texto, senao o
  Excel come o zero -- e a politica do projeto proibe exportacao que deforme
  identificador;
- **streaming**: cursor nomeado, e a auditoria fechando com a contagem REAL;
- **auditoria de download que falhou**: o rastro tem que sobreviver ao erro.
"""

import csv
import io
import os
from decimal import Decimal
from pathlib import Path

import psycopg2
import pytest
from alembic import command

from backend import migracao
from catering import auditoria, contrato
from catering.carga import carregar_tudo
from catering.carga import dimensoes
from catering.carga.fonte_csv import FonteCSV
from catering.consulta import download, matriz, planilha, recorte
from tests.conftest import consultar
from tests.test_catering_matriz import _semear_entrada, _semear_saida

DIRETORIO_DW = Path(__file__).resolve().parent.parent / "docs" / "Analise"

tem_extracao = pytest.mark.skipif(
    not (DIRETORIO_DW / "dm_volumetriaRecebimento.csv").exists(),
    reason="docs/Analise/ e gitignored -- roda onde a extracao de 21/ago existe",
)


def _conn():
    return psycopg2.connect(os.environ["DATABASE_URL"])


# ============================================================== schema
def test_migration_0021_cria_a_auditoria_e_volta(banco_migrado):
    colunas = {
        linha[0]: (linha[1], linha[2] == "YES")
        for linha in consultar(
            "SELECT column_name, data_type, is_nullable FROM information_schema.columns"
            " WHERE table_schema='public' AND table_name='cat_auditoria'"
        )
    }
    assert colunas, "cat_auditoria nao foi criada"
    # `usuario` nulavel de proposito: login e o V3.4, e nao se inventa ator
    assert colunas["usuario"][1] is True, \
        "usuario tem que aceitar nulo -- login e o V3.4"
    assert colunas["recorte"][0] == "jsonb"
    assert colunas["evento"][1] is False

    # os CHECK recusam valor desconhecido
    conn = _conn()
    try:
        with conn.cursor() as cur:
            with pytest.raises(psycopg2.Error):
                cur.execute("INSERT INTO cat_auditoria (evento) VALUES ('consulta')")
    finally:
        conn.rollback()
        conn.close()

    command.downgrade(migracao._config(), "0020_cat_cargas_fonte")
    assert not consultar(
        "SELECT 1 FROM information_schema.tables"
        " WHERE table_schema='public' AND table_name='cat_auditoria'"
    )
    command.upgrade(migracao._config(), "head")


# ============================================================ planilha
def test_planilha_pagina_com_ordem_total(cursor):
    """Sem ordem total a paginacao mente: pagina 2 repete linha da 1 e omite
    outra, sem erro nenhum. A ordenacao termina na chave natural, que e unica."""
    for i in range(1, 6):
        _semear_entrada(cursor, gem=f"{i:010d}", peso=f"{i}.000")

    vistas, paginas = [], []
    for pagina in (1, 2, 3):
        resultado = planilha.planilha(cursor, recorte.Filtros(
            de="2026-01", ate="2026-01", pagina=pagina))
        paginas.append(resultado)
        vistas += [l["guia"] for l in resultado["linhas"]]

    assert paginas[0]["paginacao"]["total_linhas"] == 5
    assert len(set(vistas)) == len(vistas), "a paginacao repetiu linha"

    # a ordenacao tem que ser estavel entre execucoes
    de_novo = planilha.planilha(cursor, recorte.Filtros(
        de="2026-01", ate="2026-01", pagina=1))
    assert [l["guia"] for l in de_novo["linhas"]] == \
        [l["guia"] for l in paginas[0]["linhas"]]


def test_planilha_e_estreita_na_tela(cursor):
    """A tela mostra a lente escolhida, nao as 16 medidas da expedicao -- seria
    o 'indo pro lado' que a Matriz evitou, e a planilha e mais larga que ela."""
    _semear_entrada(cursor)
    entrada = planilha.planilha(cursor, recorte.Filtros(
        de="2026-01", ate="2026-01", movimento="rec", lente="liq"))
    assert [c["chave"] for c in entrada["colunas"]] == [
        "dia", "unidade", "cliente", "guia", "operacao", "tipo_estoque", "valor"
    ]

    _semear_saida(cursor)
    saida = planilha.planilha(cursor, recorte.Filtros(
        de="2026-01", ate="2026-01", movimento="exp", lente="liq"))
    # uma coluna por faixa: a mesma medida em tres estados
    assert [c["chave"] for c in saida["colunas"]][-3:] == list(contrato.FAIXAS)
    linha = saida["linhas"][0]
    assert linha["solicitado"] == Decimal("100.000")
    assert linha["atendido"] == Decimal("80.000")
    assert linha["separado"] == Decimal("70.000")
    assert linha["guia"], "a guia aparece na planilha (e nao na Matriz)"

    # pallet na saida: contexto sem coluna de numero, e nao erro
    pallet = planilha.planilha(cursor, recorte.Filtros(
        de="2026-01", ate="2026-01", movimento="exp", lente="pal"))
    assert [c["chave"] for c in pallet["colunas"]] == [
        "dia", "unidade", "cliente", "guia", "operacao", "tipo_estoque"
    ]


def test_pagina_alem_do_fim_avisa_em_vez_de_quebrar(cursor):
    _semear_entrada(cursor)
    resultado = planilha.planilha(cursor, recorte.Filtros(
        de="2026-01", ate="2026-01", pagina=99))
    assert resultado["linhas"] == []
    assert any("além do fim" in a for a in resultado["avisos"])


@tem_extracao
def test_planilha_somada_bate_com_a_matriz(banco_migrado):
    """**O aceite do V3.3.** Detalhe e agregacao contam as mesmas linhas.

    Recorte estreito de propósito (uma unidade, um mes): o aceite e sobre a
    igualdade, e varrer 787 paginas nao provaria mais nada."""
    carregar_tudo(FonteCSV(DIRETORIO_DW))
    conn = _conn()
    try:
        with conn.cursor() as cur:
            base = dict(de="2026-03", ate="2026-03", movimento="rec",
                        lente="liq", unidades=("CWBIII",))

            resultado = matriz.matriz(cur, recorte.Filtros(**base))
            total_matriz = resultado["total"]["2026-03"]

            somado, pagina, linhas_vistas = Decimal(0), 1, 0
            while True:
                pag = planilha.planilha(cur, recorte.Filtros(**base, pagina=pagina))
                if not pag["linhas"]:
                    break
                for linha in pag["linhas"]:
                    if linha["valor"] is not None:
                        somado += linha["valor"]
                linhas_vistas += len(pag["linhas"])
                pagina += 1

            assert linhas_vistas == pag["paginacao"]["total_linhas"], \
                "a paginacao perdeu ou repetiu linha"
            assert somado == total_matriz, (
                f"planilha somada {somado} != total da Matriz {total_matriz} "
                "-- detalhe e agregacao discordam"
            )
    finally:
        conn.close()


# ============================================================ download
def _ler_csv(pedacos):
    texto = "".join(pedacos)
    assert texto.startswith("﻿"), \
        "sem BOM o Excel estraga os acentos no duplo clique"
    return list(csv.reader(io.StringIO(texto[1:]), delimiter=";"))


def test_csv_sai_no_formato_do_excel(cursor):
    """`;`, BOM, decimal com virgula e data DD/MM/AAAA -- e assim que o arquivo
    vai ser aberto."""
    _semear_entrada(cursor, peso="1234.567", calendario="2026-01-05")
    cursor.connection.commit()

    linhas = _ler_csv(download.gerar_csv(recorte.Filtros(de="2026-01", ate="2026-01")))
    cabecalho, dados = linhas[0], linhas[1]
    assert cabecalho[:4] == ["Dia", "Unidade", "Cliente", "Tipo de estoque"]
    assert dados[0] == "05/01/2026", "data fora do formato brasileiro"

    coluna = {nome: i for i, nome in enumerate(cabecalho)}
    assert dados[coluna["qtde_peso2"]] == "1234,567", "decimal deveria ter virgula"
    # o valor cru do identificador, correto -- o Excel e que trunca, e a tela avisa
    assert dados[coluna["num_gem"]] == "0000000001"


def test_csv_leva_a_linha_inteira_com_procedencia(cursor):
    """Derivadas E contrato cru: e o que permite conferir 'o DW diz RMSPV, a
    tela mostra RMSPIV' sem abrir o banco.

    A dimensao precisa estar populada para o de-para valer: a sigla exibida vive
    em `cat_unidades`, e sem ela o `COALESCE` cai para a sigla da fonte -- que e
    o comportamento certo (nao ha FK, e linha nova nao pode desaparecer), com
    teste proprio em `test_catering_matriz.py`."""
    _semear_entrada(cursor, sigla="RMSPV")
    cursor.connection.commit()
    dimensoes.atualizar()

    linhas = _ler_csv(download.gerar_csv(recorte.Filtros(de="2026-01", ate="2026-01")))
    cabecalho, dados = linhas[0], linhas[1]
    coluna = {nome: i for i, nome in enumerate(cabecalho)}

    assert dados[coluna["Unidade"]] == "RMSPIV", "a sigla exibida"
    assert dados[coluna["nk_wms_filial"]] == "RMSPV", "o que o DW mandou"
    for nome, _t, _n in contrato.colunas("rec"):
        assert nome in coluna, f"o download nao leva a coluna {nome}"


def test_download_e_auditado_com_o_recorte_e_a_contagem(cursor):
    _semear_entrada(cursor, gem="0000000001")
    _semear_entrada(cursor, gem="0000000002")
    cursor.connection.commit()

    filtros = recorte.Filtros(de="2026-01", ate="2026-01", unidades=("RMSPII",))
    registro = auditoria.abrir("download", recorte=filtros.como_dict(),
                               formato="csv", ip="127.0.0.1")
    linhas = _ler_csv(download.gerar_csv(filtros, registro))
    assert len(linhas) == 3          # cabecalho + 2

    gravado = consultar(
        "SELECT evento, usuario, recorte, formato, linhas, status, ip"
        " FROM cat_auditoria WHERE id = %s", (registro,)
    )[0]
    evento, usuario, recorte_gravado, formato, contagem, status, ip = gravado
    assert evento == "download"
    assert usuario is None, "login e o V3.4 -- nao se inventa ator"
    assert recorte_gravado["unidades"] == ["RMSPII"], \
        "a auditoria tem que dizer exatamente qual recorte saiu"
    assert (formato, contagem, status, ip) == ("csv", 2, "ok", "127.0.0.1")


def test_download_que_falha_deixa_rastro_de_falha(cursor, monkeypatch):
    """Download interrompido nao pode aparecer como concluido -- e o caso que
    mais interessa numa auditoria."""
    _semear_entrada(cursor)
    cursor.connection.commit()

    def sql_quebrado(_filtros):
        return "SELECT coluna_que_nao_existe FROM cat_fato_recebimento", {}

    monkeypatch.setattr(download, "_sql", sql_quebrado)
    filtros = recorte.Filtros(de="2026-01", ate="2026-01")
    registro = auditoria.abrir("download", recorte=filtros.como_dict(), formato="csv")

    with pytest.raises(psycopg2.Error):
        list(download.gerar_csv(filtros, registro))

    status, erro = consultar(
        "SELECT status, erro FROM cat_auditoria WHERE id = %s", (registro,))[0]
    assert status == "erro"
    assert erro, "a auditoria da falha ficou sem mensagem"


def test_xlsx_protege_o_zero_a_esquerda(cursor):
    """O CSV nao consegue impedir o Excel de comer o zero a esquerda; o xlsx
    consegue, escrevendo a coluna como texto. A lista de colunas protegidas sai
    do contrato, nao da minha memoria."""
    from openpyxl import load_workbook

    _semear_entrada(cursor, gem="0000000609")
    cursor.connection.commit()

    conteudo = download.gerar_xlsx(recorte.Filtros(de="2026-01", ate="2026-01"))
    livro = load_workbook(io.BytesIO(conteudo))
    aba = livro["volumetria"]
    linhas = list(aba.values)
    coluna = {nome: i for i, nome in enumerate(linhas[0])}

    for nome in contrato.IDENTIFICADORES_TEXTO:
        valor = linhas[1][coluna[nome]]
        assert isinstance(valor, str), \
            f"{nome} saiu como {type(valor).__name__} -- o zero a esquerda morreu"
    assert linhas[1][coluna["num_gem"]] == "0000000609"
    # medida continua numero, para o Excel poder somar
    assert not isinstance(linhas[1][coluna["qtde_peso2"]], str)


def test_xlsx_recusa_acima_do_teto(cursor, monkeypatch):
    """xlsx nao streama. Acima do teto a mensagem manda para o CSV, em vez de o
    servidor morrer sem explicacao."""
    _semear_entrada(cursor)
    cursor.connection.commit()
    monkeypatch.setattr(download, "TETO_XLSX", 0)

    with pytest.raises(download.DownloadGrandeDemais, match="CSV"):
        download.gerar_xlsx(recorte.Filtros(de="2026-01", ate="2026-01"))


@tem_extracao
def test_download_bate_com_o_csv_de_origem(banco_migrado):
    """O arquivo baixado somado contra o CSV do DW, no mesmo recorte."""
    carregar_tudo(FonteCSV(DIRETORIO_DW))
    filtros = recorte.Filtros(de="2026-02", ate="2026-02", movimento="rec",
                              unidades=("CWBIII",))

    baixado = _ler_csv(download.gerar_csv(filtros))
    coluna = {nome: i for i, nome in enumerate(baixado[0])}
    somado = Decimal(0)
    for linha in baixado[1:]:
        bruto = linha[coluna["qtde_peso2"]]
        if bruto:
            somado += Decimal(bruto.replace(",", "."))

    # o mesmo recorte, lido direto da fonte
    esperado = Decimal(0)
    vistas = 0
    caminho = DIRETORIO_DW / "dm_volumetriaRecebimento.csv"
    with caminho.open(encoding="utf-8", newline="") as arquivo:
        for linha in csv.DictReader(arquivo, delimiter=";"):
            if linha["NK_CALENDARIO"][:7] != "2026-02":
                continue
            if linha["NK_WMS_FILIAL"].strip() != "CWBIII":
                continue
            vistas += 1
            valor = linha["QTDE_PESO2"].strip()
            if valor:
                esperado += Decimal(valor)

    assert len(baixado) - 1 == vistas, "o download levou outra quantidade de linhas"
    assert somado == esperado, f"download {somado} != fonte {esperado}"
